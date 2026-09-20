import io
import os
from core.logger import logger


def extrair_texto_de_arquivo(file_bytes: bytes, filename: str) -> str:
    """
    Extrai o conteúdo de texto de arquivos nos formatos:
    - PDF (.pdf) via pypdf
    - Word (.docx) via python-docx
    - Texto puro (.txt)
    """
    if not file_bytes:
        raise ValueError("O arquivo enviado está vazio.")

    ext = os.path.splitext(filename)[1].lower() if filename else ""
    logger.info(f"[EXTRAIR TEXTO] Processando arquivo '{filename}' (extensão: {ext}, tamanho: {len(file_bytes)} bytes)")

    if ext == ".pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            text_parts = []
            for page in reader.pages:
                page_text = page.extract_text() or ""
                if page_text.strip():
                    text_parts.append(page_text.strip())
            extracted = "\n\n".join(text_parts)
        except Exception as exc:
            logger.error(f"[EXTRAIR TEXTO] Erro ao ler PDF '{filename}': {exc}")
            raise ValueError(f"Não foi possível extrair o texto do arquivo PDF: {exc}")

    elif ext in [".docx", ".doc"]:
        if ext == ".doc":
            try:
                import docx
                doc = docx.Document(io.BytesIO(file_bytes))
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                extracted = "\n\n".join(paragraphs)
            except Exception:
                raise ValueError(
                    "O arquivo está no formato Word 97-2003 (.doc legado). "
                    "Por favor, salve ou converta o arquivo como .docx ou .pdf e tente novamente."
                )
        else:
            try:
                import docx
                doc = docx.Document(io.BytesIO(file_bytes))
                paragraphs = []
                for p in doc.paragraphs:
                    if p.text.strip():
                        paragraphs.append(p.text.strip())
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                        if row_text:
                            paragraphs.append(row_text)
                extracted = "\n\n".join(paragraphs)
            except Exception as exc:
                logger.error(f"[EXTRAIR TEXTO] Erro ao ler DOCX '{filename}': {exc}")
                raise ValueError(f"Não foi possível extrair o texto do arquivo DOCX: {exc}")

    elif ext == ".txt":
        try:
            extracted = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                extracted = file_bytes.decode("latin-1")
            except Exception as exc:
                logger.error(f"[EXTRAIR TEXTO] Erro ao decodificar TXT '{filename}': {exc}")
                raise ValueError("Não foi possível decodificar o arquivo de texto. Use codificação UTF-8.")

    else:
        raise ValueError(
            f"Formato de arquivo '{ext or 'desconhecido'}' não suportado. "
            "Por favor, envie arquivos no formato .pdf, .docx ou .txt."
        )

    extracted = extracted.strip()
    if not extracted:
        raise ValueError(
            "Nenhum texto legível foi encontrado no arquivo enviado. "
            "Verifique se o arquivo não é uma imagem escaneada sem OCR."
        )

    return extracted
