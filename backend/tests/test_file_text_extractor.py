import io
import pytest
from fastapi.testclient import TestClient
from unittest.mock import patch, MagicMock
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
import docx

import models
from database import Base, get_db
from main import app
import security
from services.file_extractor import extrair_texto_de_arquivo

# In-memory SQLite for testing
SQLALCHEMY_DATABASE_URL = "sqlite:///:memory:"
engine = create_engine(
    SQLALCHEMY_DATABASE_URL,
    connect_args={"check_same_thread": False},
    poolclass=StaticPool
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)


def override_get_db():
    db = TestingSessionLocal()
    try:
        yield db
    finally:
        db.close()


@pytest.fixture(autouse=True)
def setup_database():
    Base.metadata.create_all(bind=engine)
    app.dependency_overrides[get_db] = override_get_db
    app.dependency_overrides[security.get_api_key] = lambda: "test-api-key"
    yield
    Base.metadata.drop_all(bind=engine)
    app.dependency_overrides.clear()


client = TestClient(app)


def test_extrair_texto_txt():
    """Deve extrair o texto de arquivos .txt em UTF-8."""
    conteudo = "Dia 1 às 09:00 - Olá pessoal! Sejam bem-vindos ao grupo VIP."
    file_bytes = conteudo.encode("utf-8")
    resultado = extrair_texto_de_arquivo(file_bytes, "roteiro.txt")
    assert resultado == conteudo


def test_extrair_texto_txt_latin1():
    """Deve extrair texto de arquivos .txt codificados em Latin-1 com caracteres acentuados."""
    conteudo = "Atenção: Promoção de Lançamento às 19:00!"
    file_bytes = conteudo.encode("latin-1")
    resultado = extrair_texto_de_arquivo(file_bytes, "roteiro_latin1.txt")
    assert "Atenção" in resultado


def test_extrair_texto_docx():
    """Deve extrair parágrafos e tabelas de um documento .docx real."""
    doc = docx.Document()
    doc.add_paragraph("Dia 1 às 10:00 - Texto: Primeiro aviso do lançamento.")
    doc.add_paragraph("Dia 2 às 15:00 - Enquete: Qual seu interesse?")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Horário: 20:00"
    table.rows[0].cells[1].text = "Live Liberada!"

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    resultado = extrair_texto_de_arquivo(docx_bytes, "campanha.docx")
    assert "Primeiro aviso do lançamento." in resultado
    assert "Qual seu interesse?" in resultado
    assert "Live Liberada!" in resultado


def test_extrair_texto_pdf_mock():
    """Deve extrair texto de páginas de PDF."""
    with patch("pypdf.PdfReader") as mock_reader_class:
        mock_reader = MagicMock()
        mock_page1 = MagicMock()
        mock_page1.extract_text.return_value = "Página 1: Boas vindas às 09h"
        mock_page2 = MagicMock()
        mock_page2.extract_text.return_value = "Página 2: Oferta imperdível às 20h"
        mock_reader.pages = [mock_page1, mock_page2]
        mock_reader_class.return_value = mock_reader

        resultado = extrair_texto_de_arquivo(b"%PDF-1.4...", "roteiro.pdf")
        assert "Página 1: Boas vindas às 09h" in resultado
        assert "Página 2: Oferta imperdível às 20h" in resultado


def test_extrair_texto_formato_invalido():
    """Deve lançar ValueError para extensões não suportadas."""
    with pytest.raises(ValueError, match="Formato de arquivo.*não suportado"):
        extrair_texto_de_arquivo(b"conteudo binario", "imagem.png")


def test_extrair_texto_arquivo_vazio():
    """Deve lançar ValueError para arquivos sem conteúdo."""
    with pytest.raises(ValueError, match="O arquivo enviado está vazio"):
        extrair_texto_de_arquivo(b"", "roteiro.txt")


def test_endpoint_extrair_texto_arquivo_sucesso():
    """Endpoint POST /mensagens/extrair-texto-arquivo deve retornar o texto extraído com status 200."""
    conteudo = "Dia 3 às 18:00 - Áudio: Ouça este recado importante."
    files = {
        "file": ("roteiro_campanha.txt", io.BytesIO(conteudo.encode("utf-8")), "text/plain")
    }

    response = client.post("/mensagens/extrair-texto-arquivo", files=files, headers={"X-API-Key": "test-api-key"})
    assert response.status_code == 200
    data = response.json()
    assert data["filename"] == "roteiro_campanha.txt"
    assert data["texto"] == conteudo


def test_endpoint_extrair_texto_arquivo_erro_formato():
    """Endpoint deve retornar 400 se o arquivo for de extensão inválida."""
    files = {
        "file": ("malicioso.exe", io.BytesIO(b"MZ\x90\x00"), "application/octet-stream")
    }

    response = client.post("/mensagens/extrair-texto-arquivo", files=files, headers={"X-API-Key": "test-api-key"})
    assert response.status_code == 400
    assert "não suportado" in response.json()["detail"]
